from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

content = r'''# Hướng dẫn xây dựng và bảo vệ project Parking YOLO
Nhận diện trạng thái ô đỗ xe bằng Python và mô hình YOLO có sẵn
Tài liệu học và chuẩn bị trả lời giáo viên • Ngày 26 tháng 10 năm 2026
## Mục tiêu cần hiểu
Project đọc video bãi xe hoặc camera góc cố định, dùng YOLO phát hiện ô tô, đối chiếu khung xe với các vùng ô đỗ đã vẽ và xác nhận trạng thái theo thời gian. Kết quả gồm ô có xe, ô còn trống và ô đang kiểm tra.
Phần AI sử dụng mô hình pretrained của Ultralytics. Phần ứng dụng gồm công cụ vẽ ô, xử lý hình học, xác nhận trạng thái và giao diện. Project hiện chưa có quy trình tự huấn luyện mô hình.
## Cách sử dụng tài liệu
Đọc các bước theo thứ tự để hiểu cách xây dựng lại hệ thống. Mỗi phần chỉ rõ tên tệp, hàm quan trọng, cách xử lý và những ý có thể giải thích với giáo viên. Cuối tài liệu có câu hỏi thường gặp và bài giới thiệu mẫu.
## Nội dung
1. Xác định bài toán và công nghệ
2. Chuẩn bị môi trường và cấu trúc project
3. Chuẩn bị dữ liệu và chọn mô hình
4. Vẽ và lưu vùng ô đỗ
5. Phát hiện ô tô bằng YOLO
6. Xác định trạng thái bằng diện tích giao
7. Xác nhận trạng thái theo thời gian
8. Giao diện và danh sách camera
9. Chạy và kiểm thử project
10. Câu hỏi về dữ liệu và AI
11. Câu hỏi về code và giới hạn
12. Bài nói mẫu và tài liệu tham khảo
===
# Bước 1 Xác định bài toán và chia chức năng
## Đầu vào và đầu ra
Đầu vào: video hoặc camera, trọng số YOLO và file JSON chứa tọa độ ô đỗ. Camera cần giữ cố định để các tọa độ đã vẽ còn đúng. Đầu ra: hình có khung xe, viền từng ô và số lượng ô có xe, trống, đang kiểm tra.
## Quy trình xử lý
Đọc một khung hình → phát hiện ô tô → lấy khung bao xe → tính phần giao với từng ô → theo dõi kết quả theo thời gian → tô màu và hiển thị thống kê. Lặp đến khi hết video hoặc người dùng dừng.
| Công nghệ | Vai trò trong project |
| Python | Ngôn ngữ chính kết nối các phần xử lý. |
| Ultralytics YOLO | Nạp trọng số và dự đoán khung bao xe. |
| OpenCV | Đọc nguồn, xử lý chuột, vẽ và tính diện tích giao. |
| NumPy | Biểu diễn tọa độ đa giác bằng mảng số. |
| Tkinter và ttk | Tạo giao diện desktop chọn nguồn và thiết lập. |
| JSON | Lưu kích thước ảnh và tọa độ các vùng ô đỗ. |
## Vì sao chọn hướng này
Bãi xe có các ô cố định nên có thể đánh dấu thủ công. Mô hình pretrained đã học nhận diện ô tô, giúp làm bản mẫu mà chưa cần gán nhãn hàng nghìn ảnh để train riêng. Phần hình học chuyển kết quả phát hiện xe thành trạng thái từng ô.
## Cách trả lời ngắn
“Em chia hệ thống thành nhận diện ô tô bằng YOLO, xác định xe có phủ ô đỗ và giao diện điều khiển. Vì góc quay cố định nên em lưu vùng ô một lần để dùng lại.”
===
# Bước 2 Chuẩn bị môi trường và cấu trúc project
## Tạo môi trường Python
Cài Python và VS Code, mở thư mục o_do_xe_YOLO. Có thể dùng Python 3.11 hoặc 3.12 với các gói phù hợp. Môi trường ảo giúp tách thư viện project khỏi các project khác. Chạy trong Terminal tại thư mục project:
`python -m venv .venv
`.venv\Scripts\python.exe -m pip install -r requirements.txt
`.venv\Scripts\python.exe app.py
requirements.txt khai báo ultralytics>=8.3.0,<9; opencv-python>=4.8,<5; numpy>=1.24,<3. Cần dùng bản Ultralytics thực tế hỗ trợ YOLO26; không phải mọi bản từ 8.3.0 đều hỗ trợ mô hình này. Tkinter thường đi kèm bản Python cài từ python.org.
## Chia tệp theo nhiệm vụ
| Tệp | Vai trò |
| app.py | Giao diện và khởi chạy các script xử lý. |
| main.py | Nhận diện xe, xác định và hiển thị trạng thái ô. |
| parking_draw.py | Công cụ vẽ, sửa và lưu đa giác ô đỗ. |
| parking_slots.json | Kích thước ảnh và danh sách vùng ô đã vẽ. |
| yolo26n.pt | Trọng số mô hình đã huấn luyện sẵn. |
| baixe.mp4 | Video đầu vào hiện có trong thư mục. |
| requirements.txt | Danh sách thư viện để cài môi trường. |
## Điều cần nhớ
.pt chứa tham số đã học của mô hình, không phải bộ ảnh huấn luyện. JSON chứa vùng ô đỗ, không phải nhãn huấn luyện YOLO. Có GPU chưa đồng nghĩa đang chạy bằng GPU; còn phụ thuộc PyTorch và môi trường CUDA thực tế.
===
# Bước 3 Chuẩn bị dữ liệu và chọn mô hình
## Ba loại dữ liệu cần phân biệt
Dữ liệu huấn luyện có sẵn: trọng số phát hiện đối tượng YOLO26 của Ultralytics được cung cấp pretrained trên COCO. Tác vụ phát hiện của COCO có 80 nhóm đối tượng, gồm car. Project dùng kiến thức đã học trong trọng số, chưa tự train trên COCO [1][2][3].
Video kiểm thử: baixe.mp4 hoặc luồng camera là đầu vào chạy thử. Cần ghi đúng nguồn video thực tế trong báo cáo: tự quay hay tải từ đâu, tác giả hoặc URL và ngày truy cập nếu có. Không ghi video do mình quay nếu điều đó chưa đúng.
Cấu hình ô đỗ: người dùng tự đánh dấu các góc ô. Công cụ ghi tọa độ vào JSON để chương trình biết vị trí P1, P2… trên ảnh. Đây là cấu hình không gian của bãi xe.
## Nano và Small
| Mô hình | Đặc điểm | Cách lựa chọn |
| yolo26n.pt | Nano có kích thước nhỏ hơn. | Dùng khi muốn giảm chi phí tính toán. |
| yolo26s.pt | Small lớn hơn bản Nano. | Thử khi cần cải thiện phát hiện. |
Mô hình lớn thường có chất lượng benchmark tốt hơn nhưng không bảo đảm mọi video đều tốt hơn. Nên so trên cùng video và tham số. Không tự khẳng định độ chính xác hoặc FPS của project khi chưa đo [3].
## Cách giải thích với giáo viên
“Em dùng YOLO26 pretrained trên COCO để nhận diện ô tô, chưa tự huấn luyện. Video là dữ liệu đầu vào kiểm thử, còn vùng ô đỗ được cấu hình bằng công cụ vẽ.”
Chạy suy luận bằng trọng số có sẵn không cần tải toàn bộ COCO. Chỉ khi tự huấn luyện hoặc đánh giá bằng dữ liệu có nhãn mới cần chuẩn bị bộ ảnh và nhãn tương ứng. Vẽ vùng ô trong JSON không phải fine-tuning.
===
# Bước 4 Xây dựng công cụ vẽ và lưu ô đỗ
## Mở nguồn và lấy khung hình
Trong parking_draw.py, chương trình thử đọc ảnh bằng cv2.imread. Nếu không phải ảnh, mở video hoặc camera bằng cv2.VideoCapture, lấy một khung hình làm nền. --frame cho phép chọn khung video khác bằng dòng lệnh; app hiện dùng mặc định.
## Thu thập tọa độ bằng chuột
Hàm mouse thêm điểm khi nhấn chuột trái, bỏ điểm cuối khi nhấn chuột phải. Chọn các góc tuần tự quanh ô, thường là bốn góc. Code cho phép đa giác lồi từ ba đỉnh, kiểm tra bằng cv2.isContourConvex, tránh hình lõm hoặc nối cắt chéo.
Ảnh được thu nhỏ để vừa màn hình vẽ. Tọa độ chuột phải chia cho sx, sy để đổi về ảnh gốc trước khi lưu. Nếu không quy đổi, vùng đã vẽ sẽ sai vị trí khi chạy trên ảnh gốc.
| Thao tác | Chức năng |
| Enter | Thêm ô hợp lệ và đặt tên P1, P2… |
| U hoặc chuột phải | Bỏ điểm cuối. |
| C | Hủy các điểm của ô đang vẽ. |
| D | Xóa ô đã thêm cuối cùng. |
| S | Lưu các ô đã xác nhận vào JSON. |
| Q hoặc Esc | Thoát, không tự lưu thay đổi. |
## Cấu trúc JSON minh họa
`{"image_size": [1920, 1080], "slots": [
`  {"id": "P1", "points": [[350,759], [495,753],
`                            [530,824], [362,832]]}
`]}
image_size giúp kiểm tra độ phân giải. Khi đổi góc quay dù giữ nguyên độ phân giải, code không tự nhận ra mọi thay đổi; cần vẽ lại. Nên vẽ các ô không chồng nhau và nhớ Enter trước S.
===
# Bước 5 Nạp YOLO và phát hiện ô tô
## Kiểm tra đầu vào và nạp trọng số
main.py dùng argparse nhận nguồn, JSON, model và các ngưỡng. Sau khi đọc JSON, chương trình gọi YOLO(args.model). Nếu thiếu cấu hình hoặc chưa có ô nào đã lưu, chương trình dừng để người dùng vẽ trước.
## Lọc đúng lớp ô tô
`model = YOLO(args.model)
`car_ids = [class_id
`    for class_id, name in model.names.items()
`    if name.lower() == "car"]
Code tìm ID từ tên lớp thay vì giả định mọi model có cùng thứ tự lớp. Model tự chọn vẫn phải có lớp tên car. Bản hiện tại chỉ lọc ô tô, chưa tự tính cả xe máy, xe buýt hay xe tải.
## Suy luận từng khung hình
`result = model.predict(
`    frame, classes=car_ids, conf=args.confidence,
`    imgsz=1280, verbose=False
`)[0]
`boxes = result.boxes.xyxy.cpu().tolist()
classes giới hạn đối tượng cần tìm. conf là ngưỡng lọc dự đoán, mặc định 0.15. imgsz=1280 đặt kích thước xử lý suy luận. boxes.xyxy chứa tọa độ x1, y1, x2, y2 theo ảnh đầu vào. cpu() chuyển tensor về CPU trước khi đổi thành danh sách Python.
## Hiểu đúng kết quả
Khung bao là hình chữ nhật bao quanh xe, chưa phải đường viền chính xác của thân xe. Vì thế khung có thể phủ sang ô bên cạnh khi camera nhìn xiên.
Giảm conf thường giữ thêm dự đoán, có thể giảm bỏ sót nhưng tăng nhận nhầm. Tăng conf lọc chặt hơn nhưng có thể mất xe khó nhìn. conf=0.15 không có nghĩa hệ thống chính xác 15%; cần dữ liệu có nhãn để đánh giá chất lượng.
===
# Bước 6 Tính tỉ lệ diện tích ô bị xe phủ
## Hàm occupied_slots trong main.py
Hàm duyệt từng ô, tính diện tích đa giác bằng cv2.contourArea rồi so với từng khung xe. Khung bao được đổi thành đa giác chữ nhật. cv2.intersectConvexConvex trả diện tích giao của hai vùng lồi. Chỉ cần một khung xe đạt ngưỡng là kết quả thô của ô bằng True.
`slot_area = abs(cv2.contourArea(polygon))
`overlap, _ = cv2.intersectConvexConvex(
`    polygon, rectangle
`)
`if overlap / slot_area >= threshold:
`    busy = True
## Công thức và ví dụ
Tỉ lệ phủ ô = diện tích giao giữa khung xe và ô đỗ / diện tích ô đỗ.
Ngưỡng mặc định là 0.35: khung xe cần phủ ít nhất 35% diện tích ô đã vẽ. Ví dụ P1 có diện tích 10.000 pixel vuông; phần giao là 4.000 pixel vuông. Tỉ lệ 4.000 / 10.000 = 0,40 nên đạt điều kiện có xe ở bước tính thô. Nếu giao 2.000 pixel vuông thì tỉ lệ 0,20, chưa đạt.
## Không nhầm với IoU hoặc kiểm tra tâm
IoU chia diện tích giao cho diện tích hợp. Project chia cho diện tích ô, nên gọi là tỉ lệ phủ ô. Bản hiện tại không dùng điều kiện tâm khung xe nằm trong ô. Khi trình bày, giải thích theo code đang chạy.
## Giới hạn của quy tắc
Mỗi ô được xét độc lập, chưa gán độc quyền một xe cho một ô. Khung lớn có thể làm nhiều ô cùng đạt ngưỡng. Code cũng không cộng diện tích giao của nhiều khung nhỏ thành một diện tích chung. Chọn ngưỡng cần dựa trên kiểm thử góc quay thực tế.
Trong parking_draw.py còn có hàm occupied_slots tương tự, nhưng luồng nhận diện hiện chạy hàm định nghĩa trong main.py. Khi giải thích thuật toán đang chạy, đối chiếu main.py.
===
# Bước 7 Xác nhận trạng thái theo thời gian
## Vì sao cần chờ trước khi xác nhận
YOLO có thể bỏ sót hoặc nhận nhầm ở vài khung. Nếu đổi trạng thái ngay thì ô dễ nhấp nháy. Code giữ một trạng thái ứng viên và chỉ xác nhận sau khi nó duy trì đủ thời gian.
| Biến | Ý nghĩa |
| raw_states | Kết quả thô có xe hoặc trống tại khung hiện tại. |
| candidate | Trạng thái ứng viên đang theo dõi. |
| candidate_since | Thời điểm bắt đầu theo dõi ứng viên. |
| confirmed | Trạng thái đã xác nhận; None khi chưa có kết luận. |
Khi kết quả thô đổi, candidate được đổi và bộ đếm bắt đầu lại. Khi thời gian đã qua đạt ngưỡng, confirmed nhận giá trị đó. Mặc định xác nhận có xe là 3 giây, ô trống là 2 giây.
## Ví dụ một ô chuyển trạng thái
P1 đang trống. Xe xuất hiện liên tục từ giây 10 đến giây 13: đủ 3 giây nên xác nhận có xe. Nếu kết quả bị ngắt thành trống ở giây 11 rồi lại có xe, bộ đếm có xe bắt đầu lại. Khi xe rời ô, kết quả trống phải duy trì đủ 2 giây để xác nhận.
## Camera và video dùng đồng hồ khác nhau
Camera dùng time.monotonic() đo thời gian thực. Video dùng frame_number / fps tính thời gian nội dung. Khi xử lý video chậm, đi qua 3 giây nội dung có thể mất hơn 3 giây đồng hồ. Nếu FPS nguồn không hợp lệ, code tạm dùng 30 FPS.
## Màu và thống kê
Đỏ: có xe đã xác nhận. Xanh lá: trống đã xác nhận. Vàng: chưa kết luận hoặc ứng viên khác trạng thái cũ. Trong giai đoạn chuyển trạng thái, ô được đếm vào nhóm đang kiểm tra. Tổng ba nhóm bằng số ô.
Đây là ổn định trạng thái từng ô, chưa phải tracking ID từng xe. Xe đi ngang đủ lâu vẫn có thể làm ô đạt điều kiện vì chưa có bước xác nhận xe đã đỗ hẳn.
===
# Bước 8 Xây dựng giao diện và chọn camera
## Tổ chức lớp ParkingApp
app.py tạo ba tab: Sử dụng, Thiết lập nhận diện và Thông báo. Tab đầu chọn nguồn, JSON và vẽ ô. Tab thiết lập giữ model, confidence, overlap và hai thời gian xác nhận. Tab thông báo nhận nhật ký xử lý. Bắt đầu và Dừng đặt bên dưới để dễ thao tác.
## Cách danh sách camera hoạt động
update_source_ui đổi hàng nhập video thành hàng chọn camera. scan_cameras thử các chỉ số 0 đến 20 bằng cv2.VideoCapture. Chỉ thiết bị mở được và đọc được một khung mới được liệt kê. Camera 1 tương ứng ID 0, Camera 3 tương ứng ID 2.
Nút Làm mới, trước đó có tên Tìm lại, quét lại sau khi cắm hoặc rút webcam. Danh sách này dùng chỉ số OpenCV, chưa hiển thị tên hãng và chưa tự tìm camera IP trong mạng. Camera điện thoại cần được Windows nhận như một webcam mới dùng theo cách này được.
## Luồng phụ và hàng đợi
Dò camera chạy bằng threading.Thread để giao diện không bị chặn. Luồng phụ đưa kết quả vào queue.Queue. poll_events được root.after gọi định kỳ ở luồng chính để cập nhật widget. Luồng phụ không trực tiếp sửa Tkinter.
## Gọi các script xử lý
start_editor gọi parking_draw.py. start_detection kiểm tra nguồn, file cấu hình và tham số rồi gọi main.py. launch dùng subprocess.Popen với cùng Python đang chạy app, đặt thư mục làm việc về thư mục project, đọc stdout để hiện nhật ký. Video nhận diện xuất hiện ở cửa sổ OpenCV riêng.
Khi tác vụ đang chạy, các nút chạy khác bị khóa. Nút Dừng yêu cầu kết thúc tiến trình; có bước kill nếu quá thời gian chờ. Khi đang vẽ phải nhấn S trước để giữ thay đổi. Nhấn Q trong cửa sổ OpenCV là cách kết thúc bình thường.
===
# Bước 9 Chạy thử và đánh giá project
## Thao tác trình diễn
1. Mở app.py bằng Python đã cài đủ thư viện. Chọn video hoặc camera và chờ danh sách camera nếu dùng thiết bị trực tiếp.
2. Mở JSON đúng với nguồn hoặc nhấn Tạo mới để chọn nơi lưu. Tạo mới chỉ chọn đường dẫn; file được ghi thật khi nhấn S trong công cụ vẽ.
3. Vẽ / sửa ô đỗ, Enter để thêm từng ô, S để lưu, Q để thoát. Kiểm tra model và tham số rồi nhấn Bắt đầu nhận diện.
4. Chỉ cho thầy một ô đỏ, một ô xanh, một ô vàng nếu có; giải thích tỉ lệ phủ và bộ đếm. Nhấn Q hoặc Esc để kết thúc, xem tab Thông báo nếu có lỗi.
## Lệnh tương đương
`python parking_draw.py baixe.mp4 --output parking_slots.json
`python main.py --source baixe.mp4 --slots parking_slots.json
`python parking_draw.py 0 --output camera0.json
`python main.py --source 0 --slots camera0.json
## Những tình huống cần thử
Kiểm tra xe đứng trong ô, xe rời đi, xe chỉ đi ngang, góc bị che và trường hợp chưa có camera. Đổi độ phân giải nguồn phải báo cấu hình không khớp. Đổi góc quay cần vẽ lại dù độ phân giải không đổi. Đổi 26n sang 26s không cần vẽ lại nếu nguồn giữ nguyên.
## Đo kết quả thay vì đoán số
Tự gán nhãn trạng thái đúng cho từng ô ở một số thời điểm rồi đối chiếu dự đoán. Báo rõ số mẫu, cách chọn mẫu và cách tính các ô đang kiểm tra. Tỉ lệ đúng = số quyết định đúng / tổng quyết định đánh giá; đây không phải mAP phát hiện đối tượng.
FPS xử lý = số khung xử lý / thời gian thực tế. FPS của file là tốc độ nội dung, không phải tốc độ chạy YOLO. Ghi model, máy, CPU/GPU và kích thước ảnh khi đo. Không ghi 95% chính xác hoặc thời gian thực nếu chưa có số đo.
===
# Câu hỏi thường gặp về dữ liệu và AI
## Em lấy bộ dữ liệu ở đâu
“Em dùng YOLO26 pretrained trên COCO của Ultralytics. Project chưa tự train. Video bãi xe là đầu vào kiểm thử và cần ghi nguồn thực tế; vùng ô đỗ là cấu hình em đánh dấu.”
## Em có tự huấn luyện không
“Dạ chưa. Em dùng trọng số có sẵn để suy luận. Nếu train riêng, em cần thu thập ảnh bãi xe, gán nhãn xe, chia tập train, validation và test, chạy huấn luyện rồi đánh giá trên tập giữ riêng.”
## Em làm phần nào trong project
“Phần ứng dụng tích hợp mô hình, công cụ vẽ và lưu vùng ô, tính diện tích giao, xác nhận trạng thái theo thời gian và giao diện chọn nguồn, cấu hình, chạy và dừng.” Nếu dùng AI hỗ trợ viết code, trình bày trung thực mức hỗ trợ và giải thích những phần mình đã hiểu, kiểm tra và điều chỉnh.
## YOLO là gì và AI nằm ở đâu
“YOLO là mô hình phát hiện đối tượng. Mô hình nhận ảnh và trả về khung bao xe. AI nằm ở dự đoán khung xe; diện tích giao và bộ đếm thời gian là quy tắc xử lý do chương trình định nghĩa.”
## COCO có phải bộ video bãi đỗ không
“Không. COCO là bộ ảnh gán nhãn cho các bài toán thị giác máy tính. Tác vụ phát hiện có 80 lớp, gồm ô tô. Video của project là nguồn đầu vào riêng.” [1][2]
## File pt có phải dataset không
“Không. File .pt chứa trọng số đã học. Dataset là dữ liệu dùng để học hoặc đánh giá. Nhờ trọng số, em có thể chạy suy luận mà không cần tải cả bộ ảnh huấn luyện.”
## Nano và Small khác nhau thế nào
“Nano nhỏ hơn, thường giảm chi phí tính toán. Small lớn hơn và thường đạt benchmark tốt hơn. Em cần so trên cùng video để chọn bản phù hợp máy và bãi xe.” [3]
===
# Câu hỏi thường gặp về code và giới hạn
## Vì sao phải vẽ ô trước
“Model có sẵn nhận diện xe, chưa biết các ô đỗ của bãi này. Do vị trí ô cố định nên em vẽ vùng trước, sau đó đối chiếu khung xe với vùng đó.”
## Khi nào ô được tính có xe
“Khi một khung xe phủ đủ diện tích ô, mặc định 35%, và kết quả duy trì đủ thời gian, mặc định 3 giây. Công thức là diện tích giao chia diện tích ô, không phải IoU.”
## Vì sao màu vàng và tại sao chờ
“Vàng là đang kiểm tra. Chờ giúp tránh đổi trạng thái vì vài khung hình sai. Xác nhận có xe mặc định 3 giây và ô trống 2 giây.”
## Một xe có làm hai ô cùng đỏ không
“Có thể. Code xét từng ô độc lập. Khung xe lớn có thể phủ đủ nhiều ô. Đây là hạn chế của bounding box và chưa có ràng buộc gán một xe cho một ô.”
## Camera di chuyển thì sao
“Tọa độ JSON không còn đúng, cần vẽ lại. Kiểm tra kích thước chỉ nhận ra đổi độ phân giải, chưa nhận ra mọi thay đổi góc.”
## Có biển số hoặc quản lý lượt xe chưa
“Chưa. Project tập trung vào trạng thái ô đỗ. Chưa có OCR biển số, tài khoản, thu phí, lưu lịch sử hoặc cơ sở dữ liệu lượt vào ra.”
## Nếu em nói nhận diện ổn thì chứng minh thế nào
“Em cần ghi nhãn đúng một tập mẫu rồi đối chiếu dự đoán, thống kê các lỗi và tốc độ xử lý. Việc chạy được demo chưa đủ để khẳng định độ chính xác tổng quát.”
## Hướng phát triển
Thu thập dữ liệu thực tế và đánh giá định lượng; fine-tune khi đủ nhãn; thêm tracking hoặc phân đoạn để giảm nhầm ô; lưu lịch sử và hỗ trợ camera IP bằng URL. Đây là đề xuất, chưa phải chức năng hiện có.
===
# Bài nói mẫu và tài liệu tham khảo
## Giới thiệu trong khoảng một phút
“Đề tài của em là nhận diện trạng thái ô đỗ xe bằng YOLO. Hệ thống nhận video hoặc camera cố định. Em dùng mô hình YOLO26 có sẵn để phát hiện ô tô rồi so khung bao xe với các vùng ô đỗ đã cấu hình bằng chuột.
Với từng ô, chương trình tính diện tích giao giữa khung xe và ô, chia cho diện tích ô. Nếu đạt ngưỡng và duy trì đủ thời gian, ô được xác nhận có xe. Nếu kết quả trống duy trì đủ lâu, ô được xác nhận trống. Vàng là đang kiểm tra, đỏ là có xe và xanh là trống.
Em dùng Python, Ultralytics, OpenCV, NumPy và Tkinter. Mô hình pretrained trên COCO, chưa tự train. Hạn chế chính là góc quay, chất lượng nhận diện và độ chính xác của vùng ô đã vẽ.”
## Năm ý cần nhớ
1. YOLO phát hiện xe; hình học suy ra trạng thái ô.
2. Project chưa tự train mô hình.
3. COCO, video kiểm thử và JSON ô đỗ là ba thứ khác nhau.
4. Tỉ lệ phủ ô không phải IoU.
5. Xác nhận theo thời gian giảm nhấp nháy nhưng không loại hết sai sót.
## Nguồn tham khảo chính thức
[1] COCO. Trang giới thiệu và tải dữ liệu.
https://cocodataset.org/#download
[2] Ultralytics. COCO Detection Dataset.
https://docs.ultralytics.com/datasets/detect/coco/
[3] Ultralytics. YOLO26.
https://docs.ultralytics.com/models/yolo26/
[4] OpenCV. Structural Analysis and Shape Descriptors.
https://docs.opencv.org/4.x/d3/dc0/group__imgproc__shape.html
Phần triển khai đối chiếu app.py, main.py và parking_draw.py của project hiện tại. Khi demo cần dùng đúng code và cấu hình đã kiểm tra.
'''

d=Document(); sec=d.sections[0]
sec.page_width=Cm(21); sec.page_height=Cm(29.7)
sec.top_margin=sec.bottom_margin=Cm(2)
sec.left_margin=Cm(2.5); sec.right_margin=Cm(2)
for name,size in [('Normal',12),('Title',23),('Heading 1',17),('Heading 2',13)]:
    st=d.styles[name]; st.font.name='Times New Roman'; st.font.size=Pt(size); st.font.color.rgb=RGBColor(0,0,0)
    st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Times New Roman')
    st.paragraph_format.space_after=Pt(6)
    st.paragraph_format.line_spacing=1.12
foot=sec.footer.paragraphs[0]; foot.alignment=2
foot.add_run('Parking YOLO  |  ')
f=OxmlElement('w:fldSimple'); f.set(qn('w:instr'),'PAGE'); foot._p.append(f)

def add_table(lines):
    cells=[[x.strip() for x in line.strip().strip('|').split('|')] for line in lines]
    t=d.add_table(rows=0,cols=len(cells[0])); t.autofit=False
    widths=([4.5,12] if len(cells[0])==2 else [3.5,6,7])
    for col,w in zip(t.columns,widths): col.width=Cm(w)
    bd=OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+side)
        for k,v in [('val','single'),('sz','4'),('color','D9D9D9')]: e.set(qn('w:'+k),v)
        bd.append(e)
    t._tbl.tblPr.append(bd)
    for i,row in enumerate(cells):
        tr=t.add_row()
        if i==0: tr._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
        for c,txt in zip(tr.cells,row):
            c.text=txt; c.vertical_alignment=1
            pr=c._tc.get_or_add_tcPr()
            sh=OxmlElement('w:shd'); sh.set(qn('w:fill'),'E7E6E6' if i==0 else 'FFFFFF'); pr.append(sh)
            margins=OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                e=OxmlElement('w:'+side); e.set(qn('w:w'),'75'); e.set(qn('w:type'),'dxa'); margins.append(e)
            pr.append(margins)
            for para in c.paragraphs:
                para.paragraph_format.space_after=Pt(2)
                para.paragraph_format.line_spacing=1
                for r in para.runs: r.font.size=Pt(10.5); r.bold=(i==0)
    d.add_paragraph().paragraph_format.space_after=Pt(0)

for idx,page in enumerate(content.split('\n===\n')):
    if idx: d.add_page_break()
    lines=page.splitlines(); i=0
    while i<len(lines):
        line=lines[i]
        if line.startswith('|'):
            block=[]
            while i<len(lines) and lines[i].startswith('|'): block.append(lines[i]); i+=1
            add_table(block); continue
        if line.startswith('## '): d.add_heading(line[3:],2)
        elif line.startswith('# '): d.add_heading(line[2:],0 if idx==0 else 1)
        elif line.startswith('`'):
            block=[]
            while i<len(lines) and lines[i].startswith('`'): block.append(lines[i][1:]); i+=1
            para=d.add_paragraph('\n'.join(block)); para.paragraph_format.line_spacing=1
            for r in para.runs: r.font.name='Consolas'; r.font.size=Pt(9)
            continue
        elif line:
            para=d.add_paragraph(line)
            para.paragraph_format.widow_control=True
        i+=1
out=Path(__file__).resolve().parents[1]/'Huong_dan_xay_dung_va_bao_ve_Parking_YOLO.docx'
d.core_properties.title='Hướng dẫn xây dựng và bảo vệ project Parking YOLO'
d.core_properties.author='Quốc Đạt'
d.save(out)
print(out)
print('Planned pages:',len(content.split('\n===\n')))
